#!/usr/bin/env python3
"""Build the free production documents for /araclar/ (storyboard, brief,
shoot-day checklist, consent form).

    ~/.hermes/tools/uv-0.13.5-linux-x64/uv run --with pymupdf --with python-docx \
        --with openpyxl --with pillow python scripts/belgeler/generate.py

Pipeline per document: content.py blocks -> HTML (fixed A4 pages, site fonts,
dark Rast logo) -> headless Chromium print-to-pdf -> PyMuPDF adds AcroForm
widgets at the positions Chromium reported (fillable PDF) -> preview WebP.
The same blocks also render to DOCX, Notion Markdown, and (checklist / shot
list) XLSX + CSV. Every page carries the Rast logo and the
"Rast Creative Studio · rastcreative.com" footer; files are CC BY 4.0.

Env: CHROME (default ~/.hermes/tools/chromium-1208/chrome-linux64/chrome) and
CHROME_LIBS (LD_LIBRARY_PATH for that build). Needs node + sharp for the logo.
"""
from __future__ import annotations

import base64
import csv
import html
import io
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import fitz  # PyMuPDF
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent))
import content as C  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
PUB = ROOT / "public/araclar"
DATA_OUT = ROOT / "src/data/belgeler"
HOME = Path.home()
CHROME = os.environ.get("CHROME", str(HOME / ".hermes/tools/chromium-1208/chrome-linux64/chrome"))
CHROME_LIBS = os.environ.get("CHROME_LIBS", str(HOME / ".hermes/tools/chrome-libs/root/usr/lib/x86_64-linux-gnu"))
TMP = Path(tempfile.mkdtemp(prefix="rast-belge-"))
LICENSE_URL = "https://creativecommons.org/licenses/by/4.0/deed.tr"
YEAR = "2026"

# ---------------------------------------------------------------- logo
def make_logo() -> Path:
    out = TMP / "logo-dark.png"
    js = (
        "const sharp=require('sharp');"
        f"sharp('{ROOT / 'public/logo-black.svg'}',{{density:400}}).resize(1800).trim().png().toFile('{out}')"
        ".then(()=>{}).catch(e=>{console.error(e);process.exit(1)})"
    )
    subprocess.run(["node", "-e", js], cwd=ROOT, check=True)
    im = Image.open(out)
    im.thumbnail((600, 600))
    im.save(out, optimize=True)
    return out


LOGO = make_logo()
LOGO_B64 = base64.b64encode(LOGO.read_bytes()).decode()

# ---------------------------------------------------------------- helpers
def esc(s: str) -> str:
    return html.escape(s, quote=True)


TR = str.maketrans("çğıöşüÇĞİÖŞÜâÂ", "cgiosuCGIOSUaA")


def slugify(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", s.translate(TR).lower()).strip("_")[:40] or "f"


class Names:
    def __init__(self):
        self.seen: dict[str, int] = {}

    def __call__(self, n: str) -> str:
        n = slugify(n)
        k = self.seen.get(n, 0) + 1
        self.seen[n] = k
        return n if k == 1 else f"{n}_{k}"


# ---------------------------------------------------------------- HTML
FONT_DIR = (ROOT / "public/fonts").as_uri()
LATIN = "U+0000-00FF, U+0131, U+0152-0153, U+02BB-02BC, U+02C6, U+02DA, U+02DC, U+2000-206F, U+2074, U+20AC, U+2122, U+2191, U+2193, U+2212, U+2215, U+FEFF, U+FFFD"
LATIN_EXT = "U+0100-02AF, U+0304, U+0308, U+0329, U+1E00-1E9F, U+1EF2-1EFF, U+2020, U+20A0-20AB, U+20AD-20C0, U+2113, U+2C60-2C7F, U+A720-A7FF"

CSS = f"""
@font-face{{font-family:Inter;font-weight:100 900;src:url({FONT_DIR}/inter-latin-wght-normal.woff2) format('woff2');unicode-range:{LATIN}}}
@font-face{{font-family:Inter;font-weight:100 900;src:url({FONT_DIR}/inter-latin-ext-wght-normal.woff2) format('woff2');unicode-range:{LATIN_EXT}}}
@font-face{{font-family:Grotesk;font-weight:300 700;src:url({FONT_DIR}/space-grotesk-latin-wght-normal.woff2) format('woff2');unicode-range:{LATIN}}}
@font-face{{font-family:Grotesk;font-weight:300 700;src:url({FONT_DIR}/space-grotesk-latin-ext-wght-normal.woff2) format('woff2');unicode-range:{LATIN_EXT}}}
@page{{size:A4;margin:0}}
@page land{{size:A4 landscape;margin:0}}
*{{box-sizing:border-box;margin:0;padding:0}}
html{{-webkit-print-color-adjust:exact;print-color-adjust:exact}}
body{{font-family:Inter,'DejaVu Sans',sans-serif;color:#1d1915;font-size:8.6pt;line-height:1.38;background:#fff}}
.page{{width:210mm;height:297mm;padding:11mm 13mm 16mm;position:relative;overflow:hidden;display:flex;flex-direction:column;break-after:page;page-break-after:always}}
.page:last-child{{break-after:auto;page-break-after:auto}}
.land .page{{width:297mm;height:210mm;padding:9mm 12mm 14mm;page:land}}
header{{display:flex;align-items:center;gap:5mm;padding-bottom:3mm;border-bottom:1.4pt solid #1d1915;margin-bottom:4mm}}
header img{{height:10mm;width:auto}}
header .t{{flex:1}}
header h1{{font-family:Grotesk,Inter,sans-serif;font-size:15pt;font-weight:700;letter-spacing:-.01em;line-height:1.1}}
header .sub{{font-size:7.6pt;color:#5b524a;margin-top:.8mm}}
header .meta{{text-align:right;font-size:7pt;color:#5b524a;line-height:1.35}}
header .meta b{{color:#b3501a;font-weight:600}}
main{{flex:1;min-height:0;overflow:hidden}}
footer{{position:absolute;left:13mm;right:13mm;bottom:7mm;display:flex;justify-content:space-between;align-items:center;gap:4mm;border-top:.6pt solid #cfc5b9;padding-top:1.6mm;font-size:6.8pt;color:#5b524a}}
.land footer{{left:12mm;right:12mm;bottom:6mm}}
footer b{{color:#1d1915;font-weight:600}}
footer img{{height:3.6mm;vertical-align:-0.8mm;margin-right:1.5mm}}
h2{{font-family:Grotesk,Inter,sans-serif;font-size:10.5pt;font-weight:700;color:#1d1915;margin:3.6mm 0 1.6mm;padding-left:2.2mm;border-left:2.2pt solid #b3501a;line-height:1.15}}
main>h2:first-child{{margin-top:0}}
p{{margin:1.2mm 0 1.8mm}}
.lbl{{display:block;font-size:6.6pt;font-weight:600;letter-spacing:.04em;text-transform:uppercase;color:#6b6158;margin-bottom:.5mm}}
.row{{display:flex;gap:4mm;margin:1.4mm 0 2mm}}
.row>div{{min-width:0}}
.fld{{display:block;height:6.6mm;border-bottom:.7pt solid #8f8379;background:#faf7f3}}
.area .fld{{border:.7pt solid #b5aa9f;border-radius:1.2mm;background:#faf7f3 repeating-linear-gradient(transparent 0 5.6mm,#e7dfd6 5.6mm 5.8mm)}}
.area{{margin:1.2mm 0 2.2mm}}
.checks{{margin:1.4mm 0 2.4mm}}
.opts{{display:flex;flex-wrap:wrap;gap:1.3mm 4.2mm}}
.opt{{display:inline-flex;align-items:center;gap:1.4mm;white-space:nowrap}}
.cb{{display:inline-block;width:3.5mm;height:3.5mm;border:.8pt solid #4a4139;border-radius:.6mm;background:#fff;flex:none}}
.ofld{{display:inline-block;width:34mm;height:4.6mm;border-bottom:.7pt solid #8f8379;background:#faf7f3}}
.note,.warn{{border-radius:1.6mm;padding:2mm 3mm;margin:2mm 0 2.6mm;font-size:8pt}}
.note{{background:#f4efe8;border:.6pt solid #e2d8cc}}
.warn{{background:#fff4e0;border:.9pt solid #d9902f;font-weight:600;color:#5a3a12}}
.cl{{width:100%;border-collapse:collapse;margin:.6mm 0 1.6mm}}
.cl td{{border-bottom:.5pt solid #ddd3c8;padding:.95mm 1.6mm .95mm 0;vertical-align:middle}}
.cl td.c{{width:6mm}}
.cl td.i b{{font-weight:600}}
.cl td.i small{{display:block;color:#6b6158;font-size:7pt;line-height:1.25;margin-top:.2mm}}
.cl td.n{{width:42mm;padding-right:0}}
.cl td.n .fld{{height:5.6mm}}
.grid{{width:100%;border-collapse:collapse;table-layout:fixed;margin-top:1mm}}
.grid th{{background:#1d1915;color:#fff;font-size:6.8pt;font-weight:600;text-align:left;padding:1.2mm 1.2mm;letter-spacing:.02em}}
.grid td{{border:.5pt solid #cfc5b9;padding:0}}
.grid tr:nth-child(even) td{{background:#faf7f3}}
.grid td .cell{{display:block;width:100%}}
dl{{margin:1mm 0 2mm}}
dl div{{display:grid;grid-template-columns:31mm 1fr;gap:3mm;padding:1.2mm 0;border-bottom:.5pt solid #e3dad0}}
dt{{font-weight:700;font-size:7.8pt}}
dd{{font-size:7.9pt;line-height:1.36}}
.sign{{display:flex;gap:4mm;align-items:flex-end;margin:2.4mm 0 2mm;padding:2mm 3mm;border:.6pt solid #cfc5b9;border-radius:1.6mm}}
.sign .who{{width:38mm;font-weight:700;font-size:8pt;line-height:1.25;align-self:center}}
.sign>div{{min-width:0}}
.sig{{display:block;height:12mm;border-bottom:.7pt solid #8f8379}}
ul.b{{margin:1mm 0 2mm 4.5mm}}
ul.b li{{margin:.8mm 0}}
.small{{font-size:6.9pt;color:#5b524a;line-height:1.35;margin-top:2mm}}
.ref{{display:grid;grid-template-columns:1.25fr 1fr;gap:6mm}}
.ref table{{width:100%;border-collapse:collapse;font-size:7.6pt}}
.ref th{{text-align:left;font-size:6.8pt;text-transform:uppercase;letter-spacing:.04em;color:#6b6158;border-bottom:.8pt solid #1d1915;padding:.8mm 1mm}}
.ref td{{border-bottom:.5pt solid #e3dad0;padding:.9mm 1mm;vertical-align:top}}
.ref td:first-child{{font-weight:700;white-space:nowrap}}
/* storyboard */
.sbhead{{margin-bottom:2.5mm}}
.panels{{display:grid;gap:4mm 5mm}}
.panel .fr{{position:relative;overflow:hidden;border:.9pt solid #1d1915;background:#fff}}
.panel .no{{display:flex;justify-content:space-between;align-items:flex-end;font-size:7pt;font-weight:700;margin-bottom:.8mm}}
.panel .no .fld{{display:inline-block;width:14mm;height:4.4mm}}
.thirds{{position:absolute;inset:0}}
.thirds i{{position:absolute;background:#e4dbd1}}
.thirds .v{{top:0;bottom:0;width:.5pt}}
.thirds .h{{left:0;right:0;height:.5pt}}
.safe{{position:absolute;left:4%;right:12%;top:11.5%;bottom:22%;border:.7pt dashed #3a9d68}}
.safe span{{position:absolute;top:-3.2mm;left:0;font-size:5.4pt;color:#3a9d68;white-space:nowrap}}
.mini{{display:flex;gap:2mm;margin-top:1.2mm}}
.mini>div{{flex:1;min-width:0}}
.mini .lbl{{font-size:5.9pt;margin-bottom:.2mm}}
.mini .fld{{height:5mm}}
.panel .area{{margin:1.2mm 0 0}}
.panel .area .lbl{{font-size:5.9pt}}
"""

FIELD_JS = """
<script>
if (location.search.includes('dump')) document.fonts.ready.then(() => {
  const out = [];
  document.querySelectorAll('.page').forEach((pg, pi) => {
    const pr = pg.getBoundingClientRect();
    const m = pg.querySelector('main');
    if (m.scrollHeight > m.clientHeight + 1) out.push({overflow: pi + 1, by: m.scrollHeight - m.clientHeight});
    pg.querySelectorAll('[data-f]').forEach(el => {
      const r = el.getBoundingClientRect();
      if (r.right > pr.right - 2 || r.bottom > pr.bottom - 2) out.push({clip: el.dataset.n, page: pi + 1});
      out.push({p: pi, t: el.dataset.f, n: el.dataset.n, x: r.left - pr.left, y: r.top - pr.top, w: r.width, h: r.height});
    });
  });
  const pre = document.createElement('pre'); pre.id = 'FIELDS'; pre.textContent = JSON.stringify(out);
  document.body.appendChild(pre);
});
</script>
"""


class HtmlDoc:
    def __init__(self, title: str, sub: str, slug: str, landscape=False):
        self.title, self.sub, self.slug, self.land = title, sub, slug, landscape
        self.pages: list[list[str]] = [[]]
        self.names = Names()

    def add(self, s: str):
        self.pages[-1].append(s)

    def brk(self):
        self.pages.append([])

    # fields
    def fld(self, name: str, kind="text", cls="fld", extra="") -> str:
        return f'<span class="{cls}" data-f="{kind}" data-n="{self.names(name)}"{extra}></span>'

    def cb(self, name: str) -> str:
        return f'<span class="cb" data-f="check" data-n="{self.names(name)}"></span>'

    def block(self, b):
        t = b[0]
        if t == "h":
            self.add(f"<h2>{esc(b[1])}</h2>")
        elif t == "p":
            self.add(f"<p>{esc(b[1])}</p>")
        elif t in ("note", "warn"):
            self.add(f'<div class="{t}">{esc(b[1])}</div>')
        elif t == "small":
            self.add(f'<p class="small">{esc(b[1])}</p>')
        elif t == "row":
            cells = "".join(
                f'<div style="flex:{fl}"><span class="lbl">{esc(lb)}</span>{self.fld(n)}</div>' for lb, n, fl in b[1]
            )
            self.add(f'<div class="row">{cells}</div>')
        elif t == "area":
            _, lb, n, hmm = b
            self.add(
                f'<div class="area"><span class="lbl">{esc(lb)}</span>'
                f'{self.fld(n, "area", extra=f" style=height:{hmm}mm")}</div>'
            )
        elif t == "checks":
            _, lb, n, opts = b
            items = []
            for i, o in enumerate(opts):
                extra = self.fld(f"{n}_{i+1}_txt", cls="ofld") if o.endswith(":") else ""
                items.append(f'<span class="opt">{self.cb(f"{n}_{i+1}")}{esc(o)}{extra}</span>')
            self.add(f'<div class="checks"><span class="lbl">{esc(lb)}</span><div class="opts">{"".join(items)}</div></div>')
        elif t == "checklist":
            _, n, items = b
            rows = "".join(
                f'<tr><td class="c">{self.cb(f"{n}_{i+1}")}</td><td class="i"><b>{esc(it)}</b><small>{esc(tip)}</small></td>'
                f'<td class="n">{self.fld(f"{n}_{i+1}_not")}</td></tr>'
                for i, (it, tip) in enumerate(items)
            )
            self.add(f'<table class="cl">{rows}</table>')
        elif t == "table":
            _, n, cols, nrows, rh = b
            head = "".join(f'<th style="width:{w}%">{esc(h)}</th>' for h, w in cols)
            body = ""
            for r in range(nrows):
                tds = ""
                for ci, (h, _w) in enumerate(cols):
                    kind = "check" if h == "✓" else "text"
                    if kind == "check":
                        inner = f'<span style="display:flex;justify-content:center;align-items:center;height:{rh}mm">{self.cb(f"{n}_{r+1}_ok")}</span>'
                    else:
                        inner = self.fld(f"{n}_{r+1}_{slugify(h) or 'no'}", "cell", "cell", f" style=height:{rh}mm")
                    tds += f"<td>{inner}</td>"
                body += f"<tr>{tds}</tr>"
            self.add(f'<table class="grid"><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>')
        elif t == "dl":
            items = "".join(f"<div><dt>{esc(a)}</dt><dd>{esc(d)}</dd></div>" for a, d in b[1])
            self.add(f"<dl>{items}</dl>")
        elif t == "sign":
            cells = "".join(
                f'<div style="flex:1"><span class="lbl">{esc(lb)}</span>{self.fld("imza_" + lb)}</div>' for lb in b[1]
            )
            self.add(
                f'<div class="sign">{cells}<div style="flex:1.2"><span class="lbl">İmza</span><span class="sig"></span></div></div>'
            )
        elif t == "signrow":
            _, who, labels, key = b
            cells = "".join(
                f'<div style="flex:{1.4 if i == 0 else 0.8}"><span class="lbl">{esc(lb)}</span>{self.fld(f"{key}_{lb}")}</div>'
                for i, lb in enumerate(labels)
            )
            self.add(
                f'<div class="sign"><div class="who">{esc(who)}</div>{cells}'
                f'<div style="flex:1.1"><span class="lbl">İmza</span><span class="sig"></span></div></div>'
            )
        elif t == "bullets":
            self.add('<ul class="b">' + "".join(f"<li>{esc(x)}</li>" for x in b[1]) + "</ul>")
        elif t == "ref":
            self.add(ref_html())
        elif t == "break":
            self.brk()
        elif t == "html":
            self.add(b[1])
        else:
            raise ValueError(t)

    def render(self) -> str:
        n = len(self.pages)
        logo = f"data:image/png;base64,{LOGO_B64}"
        url = f"rastcreative.com/araclar/{self.slug}/"
        out = []
        for i, body in enumerate(self.pages, 1):
            out.append(
                f'<section class="page"><header><img src="{logo}" alt="Rast Creative">'
                f'<div class="t"><h1>{esc(self.title)}</h1><div class="sub">{esc(self.sub)}</div></div>'
                f'<div class="meta"><b>Ücretsiz şablon</b><br>{url}</div></header>'
                f'<main>{"".join(body)}</main>'
                f'<footer><span><img src="{logo}" alt=""><b>Rast Creative Studio · rastcreative.com</b></span>'
                f'<span>CC BY 4.0 · atıfla serbestçe kullanın ve uyarlayın</span><span>Sayfa {i} / {n}</span></footer></section>'
            )
        cls = "land" if self.land else ""
        return (
            f'<!doctype html><html lang="tr"><head><meta charset="utf-8"><title>{esc(self.title)} · Rast Creative Studio</title>'
            f"<style>{CSS}</style></head><body class=\"{cls}\">{''.join(out)}{FIELD_JS}</body></html>"
        )


def ref_html() -> str:
    o = "".join(f"<tr><td>{esc(a)}</td><td>{esc(b)}</td><td>{esc(c)}</td></tr>" for a, b, c in C.OLCEKLER)
    h = "".join(f"<tr><td>{esc(a)}</td><td>{esc(b)}</td></tr>" for a, b in C.HAREKETLER)
    a = "".join(f"<tr><td>{esc(x)}</td><td>{esc(y)}</td></tr>" for x, y in C.ACILAR)
    return (
        '<div class="ref"><div><table><thead><tr><th>Kısaltma</th><th>Plan ölçeği</th><th>Ne için</th></tr></thead>'
        f"<tbody>{o}</tbody></table></div><div><table><thead><tr><th>Hareket</th><th>Not</th></tr></thead><tbody>{h}</tbody></table>"
        f'<table style="margin-top:3mm"><thead><tr><th>Açı</th><th>Etki</th></tr></thead><tbody>{a}</tbody></table></div></div>'
    )


# ---------------------------------------------------------------- storyboard pages
def storyboard_doc(vertical: bool) -> HtmlDoc:
    sb = C.STORYBOARD
    if vertical:
        d = HtmlDoc("Storyboard · 9:16 dikey", "Reels, TikTok, Shorts ve Story için. Yeşil kesikli çizgi yaklaşık güvenli alan.", sb["slug"], True)
    else:
        d = HtmlDoc("Storyboard · 16:9 yatay", "Tanıtım filmi, reklam ve YouTube için. Kare içindeki açık çizgiler üçler kuralı.", sb["slug"], True)
    d.block(("row", [("Proje", "proje", 3), ("Müşteri", "musteri", 2), ("Sahne", "sahne", 1), ("Tarih", "tarih", 1.4), ("Sayfa", "sayfa", 0.8)]))
    panels = []
    count = 5 if vertical else 6
    for i in range(1, count + 1):
        if vertical:
            frame = ('<div class="fr" style="width:50.6mm;height:89.95mm"><div class="thirds">'
                '<i class="v" style="left:33.333%"></i><i class="v" style="left:66.666%"></i>'
                '<i class="h" style="top:33.333%"></i><i class="h" style="top:66.666%"></i></div><div class="safe"><span>güvenli alan ≈</span></div></div>')
            below = (
                f'<div class="mini"><div><span class="lbl">Ölçek</span>{d.fld(f"p{i}_olcek")}</div>'
                f'<div><span class="lbl">Süre</span>{d.fld(f"p{i}_sure")}</div></div>'
                f'<div class="area"><span class="lbl">Aksiyon / diyalog</span>{d.fld(f"p{i}_aksiyon", "area", extra=" style=height:13mm")}</div>'
                f'<div class="mini"><div><span class="lbl">Ekran yazısı</span>{d.fld(f"p{i}_yazi")}</div></div>'
                f'<div class="mini"><div><span class="lbl">Ses / müzik</span>{d.fld(f"p{i}_ses")}</div></div>'
            )
        else:
            frame = ('<div class="fr" style="width:87.66mm;height:49.31mm"><div class="thirds">'
                '<i class="v" style="left:33.333%"></i><i class="v" style="left:66.666%"></i>'
                '<i class="h" style="top:33.333%"></i><i class="h" style="top:66.666%"></i></div></div>')
            below = (
                f'<div class="mini"><div><span class="lbl">Ölçek</span>{d.fld(f"p{i}_olcek")}</div>'
                f'<div><span class="lbl">Hareket</span>{d.fld(f"p{i}_hareket")}</div>'
                f'<div style="flex:.6"><span class="lbl">Süre</span>{d.fld(f"p{i}_sure")}</div></div>'
                f'<div class="mini"><div style="flex:2"><span class="lbl">Aksiyon / diyalog</span>{d.fld(f"p{i}_aksiyon")}</div>'
                f'<div><span class="lbl">Ses / müzik</span>{d.fld(f"p{i}_ses")}</div></div>'
            )
        panels.append(
            f'<div class="panel"><div class="no"><span>Plan {d.fld(f"p{i}_no")}</span></div>{frame}{below}</div>'
        )
    cols = 5 if vertical else 3
    d.add(f'<div class="panels" style="grid-template-columns:repeat({cols},1fr)">{"".join(panels)}</div>')
    return d


# ---------------------------------------------------------------- chrome + pdf
def chrome(args: list[str]) -> str:
    env = dict(os.environ, LD_LIBRARY_PATH=CHROME_LIBS)
    r = subprocess.run(
        [CHROME, "--headless=new", "--no-sandbox", "--disable-gpu", "--hide-scrollbars", "--allow-file-access-from-files", *args],
        env=env, capture_output=True, text=True, timeout=120,
    )
    return r.stdout


def build_pdf(doc: HtmlDoc, out: Path, meta: dict) -> dict:
    src = TMP / (out.stem + ".html")
    src.write_text(doc.render(), encoding="utf-8")
    dom = chrome(["--virtual-time-budget=8000", "--dump-dom", src.as_uri() + "?dump"])
    m = re.search(r'<pre id="FIELDS">(.*?)</pre>', dom, re.S)
    if not m:
        raise SystemExit(f"no field dump for {out.name}")
    items = json.loads(html.unescape(m.group(1)))
    problems = [x for x in items if "overflow" in x or "clip" in x]
    fields = [x for x in items if "t" in x]
    raw = TMP / (out.stem + ".raw.pdf")
    chrome(["--no-pdf-header-footer", "--virtual-time-budget=8000", f"--print-to-pdf={raw}", src.as_uri()])
    pdf = fitz.open(raw)
    if pdf.page_count != len(doc.pages):
        problems.append({"pages": pdf.page_count, "expected": len(doc.pages)})
    k = 0.75  # css px -> pt
    for f in fields:
        page = pdf[f["p"]]
        r = fitz.Rect(f["x"] * k, f["y"] * k, (f["x"] + f["w"]) * k, (f["y"] + f["h"]) * k)
        w = fitz.Widget()
        w.field_name = f["n"]
        w.rect = r
        w.border_width = 0
        w.border_color = None
        w.fill_color = None
        if f["t"] == "check":
            w.field_type = fitz.PDF_WIDGET_TYPE_CHECKBOX
            w.text_font = "ZaDb"
            w.field_value = False
            w.rect = fitz.Rect(r.x0 + 0.6, r.y0 + 0.6, r.x1 - 0.6, r.y1 - 0.6)
        else:
            w.field_type = fitz.PDF_WIDGET_TYPE_TEXT
            w.text_font = "Helv"
            w.text_fontsize = 7.5 if f["t"] == "cell" else 9
            if f["t"] == "area" or (f["t"] == "cell" and r.height > 20):
                w.field_flags |= fitz.PDF_TX_FIELD_IS_MULTILINE
        page.add_widget(w)
    pdf.set_metadata({
        "title": meta["title"],
        "author": "Rast Creative Studio",
        "subject": meta.get("subject", ""),
        "keywords": meta.get("keywords", "") + ", rastcreative.com, CC BY 4.0",
        "creator": "Rast Creative Studio · rastcreative.com",
        "producer": "rastcreative.com",
    })
    pdf.save(out, garbage=3, deflate=True)
    pdf.close()
    return {"file": out.name, "pages": len(doc.pages), "fields": len(fields), "problems": problems}


def previews(pdf_path: Path, base: str, pages: list[int], width=900) -> list[str]:
    pdf = fitz.open(pdf_path)
    names = []
    for i, p in enumerate(pages, 1):
        page = pdf[p - 1]
        zoom = width / page.rect.width
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False, annots=False)
        im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        name = f"{base}-onizleme-{i}.webp"
        im.save(pdf_path.parent / name, "WEBP", quality=80, method=6)
        names.append({"src": name, "w": im.width, "h": im.height})
        # full-size PNG for QA only (not shipped)
        im.save(TMP / f"qa-{base}-{i}.png")
    pdf.close()
    return names


# ---------------------------------------------------------------- DOCX
def docx_build(title: str, sub: str, slug: str, blocks: list, out: Path, landscape=False):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    INK = RGBColor(0x1D, 0x19, 0x15)
    AMBER = RGBColor(0xB3, 0x50, 0x1A)
    SOFT = RGBColor(0x6B, 0x61, 0x58)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    st.font.color.rgb = INK
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Mm(297), Mm(210)
    else:
        sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Mm(15))
    sec.top_margin, sec.bottom_margin = Mm(14), Mm(16)
    usable = (297 if landscape else 210) - 30

    def borders(cell, **kw):
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{edge}")
            spec = kw.get(edge)
            if spec:
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), str(spec[0]))
                el.set(qn("w:color"), spec[1])
            else:
                el.set(qn("w:val"), "nil")
            b.append(el)
        tcPr.append(b)

    def shade(cell, color):
        tcPr = cell._tc.get_or_add_tcPr()
        s = OxmlElement("w:shd")
        s.set(qn("w:val"), "clear")
        s.set(qn("w:color"), "auto")
        s.set(qn("w:fill"), color)
        tcPr.append(s)

    def row_height(row, mm):
        trPr = row._tr.get_or_add_trPr()
        h = OxmlElement("w:trHeight")
        h.set(qn("w:val"), str(int(mm * 56.7)))
        h.set(qn("w:hRule"), "atLeast")
        trPr.append(h)

    def label(par, text):
        r = par.add_run(text.upper())
        r.font.size = Pt(7)
        r.font.bold = True
        r.font.color.rgb = SOFT

    def tight(p, before=0, after=2):
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)

    # header: logo + title, footer: brand line
    hdr = sec.header
    hp = hdr.paragraphs[0]
    hp.add_run().add_picture(str(LOGO), height=Mm(8))
    hp.add_run("   " + title).bold = True
    tight(hp)
    ftr = sec.footer.paragraphs[0]
    fr = ftr.add_run(f"Rast Creative Studio · rastcreative.com   |   rastcreative.com/araclar/{slug}/   |   CC BY 4.0 — atıfla serbestçe kullanın ve uyarlayın")
    fr.font.size = Pt(7.5)
    fr.font.color.rgb = SOFT

    t = doc.add_paragraph()
    r = t.add_run(title)
    r.bold = True
    r.font.size = Pt(18)
    tight(t, 0, 0)
    s = doc.add_paragraph()
    rs = s.add_run(sub)
    rs.font.size = Pt(9)
    rs.font.color.rgb = SOFT
    tight(s, 0, 6)

    def field_table(cells: list[tuple[str, float]], h_mm=7.5):
        tb = doc.add_table(rows=2, cols=len(cells))
        tb.alignment = WD_TABLE_ALIGNMENT.LEFT
        tot = sum(w for _, w in cells)
        for i, (lb, w) in enumerate(cells):
            for rr in (0, 1):
                tb.cell(rr, i).width = Mm(usable * w / tot)
            c0 = tb.cell(0, i)
            label(c0.paragraphs[0], lb)
            tight(c0.paragraphs[0], 2, 0)
            borders(c0)
            c1 = tb.cell(1, i)
            borders(c1, bottom=(6, "8F8379"))
            shade(c1, "FAF7F3")
        row_height(tb.rows[1], h_mm)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    for b in blocks:
        k = b[0]
        if k == "h":
            p = doc.add_paragraph()
            r = p.add_run(b[1])
            r.bold = True
            r.font.size = Pt(12.5)
            r.font.color.rgb = AMBER
            tight(p, 8, 3)
        elif k == "p":
            tight(doc.add_paragraph(b[1]), 2, 4)
        elif k in ("note", "warn"):
            tb = doc.add_table(rows=1, cols=1)
            c = tb.cell(0, 0)
            c.width = Mm(usable)
            shade(c, "FFF4E0" if k == "warn" else "F4EFE8")
            borders(c, **{e: (8, "D9902F" if k == "warn" else "E2D8CC") for e in ("top", "left", "bottom", "right")})
            r = c.paragraphs[0].add_run(b[1])
            r.font.size = Pt(9)
            r.bold = k == "warn"
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif k == "small":
            p = doc.add_paragraph()
            r = p.add_run(b[1])
            r.font.size = Pt(7.5)
            r.font.color.rgb = SOFT
        elif k == "row":
            field_table([(lb, fl) for lb, _n, fl in b[1]])
        elif k == "area":
            _, lb, _n, hmm = b
            p = doc.add_paragraph()
            label(p, lb)
            tight(p, 2, 1)
            tb = doc.add_table(rows=1, cols=1)
            c = tb.cell(0, 0)
            c.width = Mm(usable)
            borders(c, **{e: (6, "B5AA9F") for e in ("top", "left", "bottom", "right")})
            shade(c, "FAF7F3")
            row_height(tb.rows[0], hmm)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif k == "checks":
            _, lb, _n, opts = b
            p = doc.add_paragraph()
            label(p, lb)
            tight(p, 3, 1)
            p = doc.add_paragraph()
            tight(p, 0, 4)
            p.add_run("   ".join("☐ " + (o + " ____________" if o.endswith(":") else o) for o in opts)).font.size = Pt(9.5)
        elif k == "checklist":
            _, _n, items = b
            tb = doc.add_table(rows=len(items), cols=3)
            for i, (it, tip) in enumerate(items):
                cs = tb.rows[i].cells
                cs[0].width, cs[1].width, cs[2].width = Mm(8), Mm(usable - 58), Mm(50)
                cs[0].paragraphs[0].add_run("☐").font.size = Pt(12)
                r = cs[1].paragraphs[0].add_run(it)
                r.bold = True
                r.font.size = Pt(9.5)
                r2 = cs[1].paragraphs[0].add_run("\n" + tip)
                r2.font.size = Pt(8)
                r2.font.color.rgb = SOFT
                for c in cs[:2]:
                    borders(c, bottom=(4, "DDD3C8"))
                borders(cs[2], bottom=(6, "8F8379"))
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif k == "table":
            _, _n, cols, nrows, rh = b
            tb = doc.add_table(rows=nrows + 1, cols=len(cols))
            tb.style = "Table Grid"
            for i, (h, w) in enumerate(cols):
                for rr in range(nrows + 1):
                    tb.cell(rr, i).width = Mm(usable * w / 100)
                c = tb.cell(0, i)
                shade(c, "1D1915")
                r = c.paragraphs[0].add_run(h)
                r.bold = True
                r.font.size = Pt(8)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            for rr in range(1, nrows + 1):
                row_height(tb.rows[rr], rh)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif k == "dl":
            tb = doc.add_table(rows=len(b[1]), cols=2)
            for i, (a, d) in enumerate(b[1]):
                c0, c1 = tb.cell(i, 0), tb.cell(i, 1)
                c0.width, c1.width = Mm(36), Mm(usable - 36)
                r = c0.paragraphs[0].add_run(a)
                r.bold = True
                r.font.size = Pt(9)
                c1.paragraphs[0].add_run(d).font.size = Pt(9)
                for c in (c0, c1):
                    borders(c, bottom=(4, "E3DAD0"))
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif k in ("sign", "signrow"):
            if k == "sign":
                who, labels = None, b[1]
            else:
                who, labels = b[1], b[2]
            cells = ([(who, 1.2)] if who else []) + [(lb, 1.2 if i == 0 else 0.8) for i, lb in enumerate(labels)] + [("İmza", 1.2)]
            tb = doc.add_table(rows=2, cols=len(cells))
            tot = sum(w for _, w in cells)
            for i, (lb, w) in enumerate(cells):
                c0, c1 = tb.cell(0, i), tb.cell(1, i)
                c0.width = c1.width = Mm(usable * w / tot)
                if who and i == 0:
                    r = c0.paragraphs[0].add_run(lb)
                    r.bold = True
                    r.font.size = Pt(9)
                    borders(c0)
                    borders(c1)
                else:
                    label(c0.paragraphs[0], lb)
                    borders(c0)
                    borders(c1, bottom=(6, "8F8379"))
            row_height(tb.rows[1], 12 if True else 8)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif k == "bullets":
            for x in b[1]:
                p = doc.add_paragraph(x, style="List Bullet")
                tight(p, 0, 2)
        elif k == "ref":
            for head, rows in (
                (("Kısaltma", "Plan ölçeği", "Ne için"), C.OLCEKLER),
                (("Hareket", "Not"), C.HAREKETLER),
                (("Açı", "Etki"), C.ACILAR),
            ):
                tb = doc.add_table(rows=len(rows) + 1, cols=len(head))
                tb.style = "Table Grid"
                for i, h in enumerate(head):
                    r = tb.cell(0, i).paragraphs[0].add_run(h)
                    r.bold = True
                    r.font.size = Pt(8.5)
                    shade(tb.cell(0, i), "F4EFE8")
                for ri, row in enumerate(rows, 1):
                    for ci, v in enumerate(row):
                        tb.cell(ri, ci).paragraphs[0].add_run(v).font.size = Pt(8.5)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif k == "break":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph()
    tight(p, 10, 0)
    r = p.add_run(
        f"Kaynak: Rast Creative Studio — rastcreative.com/araclar/{slug}/ · Lisans: CC BY 4.0 ({LICENSE_URL}). "
        "Ticari işlerde de kullanabilir ve uyarlayabilirsiniz; tek şart atıf."
    )
    r.font.size = Pt(7.5)
    r.font.color.rgb = SOFT
    cp = doc.core_properties
    cp.title, cp.author, cp.comments = title, "Rast Creative Studio", f"https://rastcreative.com/araclar/{slug}/ · CC BY 4.0"
    doc.save(out)


# ---------------------------------------------------------------- Markdown
def md_build(title: str, sub: str, slug: str, blocks: list, out: Path):
    L = [
        f"![Rast Creative Studio](https://rastcreative.com/logo-black.svg)",
        "",
        f"# {title}",
        "",
        f"_{sub}_",
        "",
    ]
    for b in blocks:
        k = b[0]
        if k == "h":
            L += [f"## {b[1]}", ""]
        elif k == "p":
            L += [b[1], ""]
        elif k == "note":
            L += [f"> 💡 {b[1]}", ""]
        elif k == "warn":
            L += [f"> ⚠️ **{b[1]}**", ""]
        elif k == "small":
            L += [f"<sub>{b[1]}</sub>", ""]
        elif k == "row":
            for lb, _n, _f in b[1]:
                L += [f"**{lb}:** ", ""]
        elif k == "area":
            L += [f"**{b[1]}:**", "", "> ", ""]
        elif k == "checks":
            L += [f"**{b[1]}**", ""] + [f"- [ ] {o + ' …' if o.endswith(':') else o}" for o in b[3]] + [""]
        elif k == "checklist":
            L += [f"- [ ] **{it}** — {tip}" for it, tip in b[2]] + [""]
        elif k == "table":
            cols = [h for h, _ in b[2]]
            L += ["| " + " | ".join(cols) + " |", "|" + "|".join("---" for _ in cols) + "|"]
            L += ["| " + " | ".join(str(i) if j == 0 else "" for j in range(len(cols))) + " |" for i in range(1, 6)]
            L += [""]
        elif k == "dl":
            L += [f"**{a}.** {d}" + "\n" for a, d in b[1]]
        elif k == "sign":
            L += [" · ".join(f"**{lb}:** ____" for lb in b[1]) + " · **İmza:** ____", ""]
        elif k == "signrow":
            L += [f"**{b[1]}** — " + " · ".join(f"{lb}: ____" for lb in b[2]) + " · İmza: ____", ""]
        elif k == "bullets":
            L += [f"- {x}" for x in b[1]] + [""]
        elif k == "ref":
            L += ["| Kısaltma | Plan ölçeği | Ne için |", "|---|---|---|"] + [f"| {a} | {x} | {c} |" for a, x, c in C.OLCEKLER] + [""]
            L += ["| Hareket | Not |", "|---|---|"] + [f"| {a} | {x} |" for a, x in C.HAREKETLER] + [""]
            L += ["| Açı | Etki |", "|---|---|"] + [f"| {a} | {x} |" for a, x in C.ACILAR] + [""]
        elif k == "break":
            L += ["---", ""]
    L += [
        "---",
        "",
        f"Kaynak: **Rast Creative Studio** — [rastcreative.com/araclar/{slug}/](https://rastcreative.com/araclar/{slug}/) · "
        f"Lisans: [CC BY 4.0]({LICENSE_URL}) (atıfla serbestçe kullanın ve uyarlayın).",
        "",
    ]
    text = "\n".join(L)
    text = re.sub(r"\n{3,}", "\n\n", text)
    out.write_text(text, encoding="utf-8")


# ---------------------------------------------------------------- XLSX / CSV
def xlsx_style(ws, title, slug, ncols):
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font

    ws.row_dimensions[1].height = 34
    img = XLImage(str(LOGO))
    img.height = 36
    img.width = int(36 * Image.open(LOGO).width / Image.open(LOGO).height)
    ws.add_image(img, "A1")
    ws.cell(1, 3, title).font = Font(bold=True, size=14)
    ws.cell(2, 3, f"Rast Creative Studio · rastcreative.com/araclar/{slug}/ · CC BY 4.0").font = Font(size=9, color="6B6158")
    ws.oddFooter.center.text = "Rast Creative Studio · rastcreative.com · CC BY 4.0"
    ws.oddFooter.right.text = "Sayfa &P / &N"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def header_row(ws, r, headers, widths):
    from openpyxl.styles import Alignment, Font, PatternFill

    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(r, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D1915")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = ws.cell(r + 1, 1)
    ws.print_title_rows = f"{r}:{r}"


def about_sheet(wb, slug, extra=""):
    from openpyxl.styles import Font

    ws = wb.create_sheet("Hakkında")
    xlsx_style(ws, "Hakkında", slug, 4)
    lines = [
        "Bu dosya Rast Creative Studio'nun ücretsiz üretim belgelerinden biridir.",
        f"Kaynak: https://rastcreative.com/araclar/{slug}/",
        f"Lisans: CC BY 4.0 ({LICENSE_URL}) — ticari işlerde de kullanabilir ve uyarlayabilirsiniz; tek şart atıf:",
        "“Şablon: Rast Creative Studio — rastcreative.com”",
        "Google E-Tablolar: Dosya > İçe aktar > Yükle ile açın; açılır listeler ve biçimler korunur.",
    ]
    if extra:
        lines.append(extra)
    for i, t in enumerate(lines, 4):
        ws.cell(i, 1, t).font = Font(size=10)
    ws.column_dimensions["A"].width = 110


def checklist_xlsx(out_xlsx: Path, out_csv: Path):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    cl = C.CHECKLIST
    wb = Workbook()
    ws = wb.active
    ws.title = "Kontrol listesi"
    xlsx_style(ws, cl["title"], cl["slug"], 6)
    ws.cell(3, 1, "Proje:").font = Font(bold=True)
    ws.cell(3, 4, "Çekim tarihi:").font = Font(bold=True)
    headers = ["Bölüm", "Madde", "İpucu", "Durum", "Sorumlu", "Not"]
    header_row(ws, 5, headers, [18, 52, 44, 11, 16, 34])
    dv = DataValidation(type="list", formula1='"✓,—,Yok,Sorun"', allow_blank=True)
    ws.add_data_validation(dv)
    r = 6
    rows = []
    for sec, items in cl["sections"]:
        for it, tip in items:
            ws.cell(r, 1, sec)
            ws.cell(r, 2, it).font = Font(bold=True)
            ws.cell(r, 3, tip).font = Font(color="6B6158")
            for c in range(1, 7):
                ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
            dv.add(ws.cell(r, 4))
            rows.append([sec, it, tip, "", "", ""])
            r += 1
    last = r - 1
    ws.auto_filter.ref = f"A5:F{last}"
    ws.conditional_formatting.add(f"A6:F{last}", FormulaRule(formula=['$D6="✓"'], fill=PatternFill("solid", fgColor="E3F4EA")))
    ws.conditional_formatting.add(f"A6:F{last}", FormulaRule(formula=['$D6="Sorun"'], fill=PatternFill("solid", fgColor="FBE3DC")))
    ws.cell(last + 2, 1, "Tamamlanan:").font = Font(bold=True)
    ws.cell(last + 2, 2, f'=COUNTIF(D6:D{last},"✓")&" / "&COUNTA(B6:B{last})')
    about_sheet(wb, cl["slug"], "Durum sütunundaki açılır listeden ✓, —, Yok ya da Sorun seçin; satır renklenir.")
    wb.properties.creator = "Rast Creative Studio"
    wb.properties.title = cl["title"]
    wb.save(out_xlsx)
    write_csv(out_csv, headers, rows)


def shotlist_xlsx(out_xlsx: Path, out_csv: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    sb = C.STORYBOARD
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan listesi"
    xlsx_style(ws, "Plan listesi (shot list)", sb["slug"], 12)
    ws.page_setup.orientation = "landscape"
    for i, lb in enumerate(["Proje:", "Çekim günü:", "Lokasyon:", "Yönetmen:"]):
        ws.cell(3, 1 + i * 3, lb).font = Font(bold=True)
    headers = [h for h, _ in C.SHOT_COLS]
    headers[-1] = "Tamam (✓)"
    header_row(ws, 5, headers, [5, 7, 7, 13, 15, 15, 16, 46, 13, 8, 16, 10])
    # vocab sheet feeds the dropdowns
    vs = wb.create_sheet("Kısaltmalar")
    xlsx_style(vs, "Plan ölçekleri, açılar ve hareketler", sb["slug"], 6)
    header_row(vs, 4, ["Kısaltma", "Plan ölçeği", "Ne için", "Açı", "Etki", "Hareket", "Not"], [14, 16, 48, 18, 34, 18, 40])
    for i, (a, b, c) in enumerate(C.OLCEKLER, 5):
        vs.cell(i, 1, a), vs.cell(i, 2, b), vs.cell(i, 3, c)
    for i, (a, b) in enumerate(C.ACILAR, 5):
        vs.cell(i, 4, a), vs.cell(i, 5, b)
    for i, (a, b) in enumerate(C.HAREKETLER, 5):
        vs.cell(i, 6, a), vs.cell(i, 7, b)
    n_o, n_a, n_h = len(C.OLCEKLER) + 4, len(C.ACILAR) + 4, len(C.HAREKETLER) + 4
    for col, ref in ((4, f"Kısaltmalar!$A$5:$A${n_o}"), (5, f"Kısaltmalar!$D$5:$D${n_a}"), (6, f"Kısaltmalar!$F$5:$F${n_h}")):
        dv = DataValidation(type="list", formula1=f"='{ref.split('!')[0]}'!{ref.split('!')[1]}", allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        L = ws.cell(6, col).column_letter
        dv.add(f"{L}6:{L}205")
    dvok = DataValidation(type="list", formula1='"✓"', allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dvok)
    dvok.add("L6:L205")
    for r in range(6, 206):
        ws.cell(r, 1, r - 5 if r < 46 else None)
        for c in range(1, 13):
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
    for i, row in enumerate(C.SHOT_EXAMPLES, 6):
        for c, v in enumerate(row, 1):
            ws.cell(i, c, v).font = Font(italic=True, color="8A7F75")
    ws.auto_filter.ref = "A5:L205"
    about_sheet(wb, sb["slug"], "Örnek satırlar (gri, italik) yalnızca yön göstermek içindir; silip kendi planlarınızı yazın.")
    wb.properties.creator = "Rast Creative Studio"
    wb.properties.title = "Plan listesi (shot list)"
    wb.save(out_xlsx)
    write_csv(out_csv, headers, C.SHOT_EXAMPLES)


def write_csv(path: Path, headers, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)


# ---------------------------------------------------------------- bundle
def readme(title, slug, files):
    return (
        f"{title}\nRast Creative Studio · rastcreative.com\n\n"
        f"Kaynak sayfa: https://rastcreative.com/araclar/{slug}/\n\n"
        "Dosyalar:\n" + "".join(f"  - {f}\n" for f in files) + "\n"
        "Nasıl kullanılır?\n"
        "  - PDF: yazdırıp elle doldurun ya da Acrobat Reader / tarayıcıda doğrudan doldurun (doldurulabilir alanlar).\n"
        "  - DOCX: Word, Google Dokümanlar ya da LibreOffice'te açıp kendinize göre düzenleyin.\n"
        "  - XLSX / CSV: Excel ya da Google E-Tablolar (Dosya > İçe aktar). CSV UTF-8'dir.\n"
        "  - MD: Notion'da Import > Markdown ile içe aktarın ya da metni yapıştırın.\n\n"
        "Lisans: CC BY 4.0 — https://creativecommons.org/licenses/by/4.0/deed.tr\n"
        "Ticari işlerde de kullanabilir, uyarlayabilir ve paylaşabilirsiniz. Tek şart atıf:\n"
        "  “Şablon: Rast Creative Studio — rastcreative.com”\n\n"
        "Bir çekimi birlikte planlamak isterseniz: https://rastcreative.com/iletisim/\n"
    )


def bundle(slug: str, title: str, files: list[Path], extra_readme: str = "") -> Path:
    z = PUB / slug / f"rast-{slug}-paketi.zip"
    names = [f.name for f in files]
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in files:
            zf.write(f, f.name)
        zf.writestr("OKU-BENI.txt", readme(title, slug, names) + extra_readme)
    return z


# ---------------------------------------------------------------- main
def planner_assets():
    """Logo for the call-sheet print view + the sample plan JSON."""
    out = PUB / "cekim-plani"
    out.mkdir(parents=True, exist_ok=True)
    im = Image.open(LOGO)
    im.thumbnail((360, 360))
    im.save(out / "rast-logo-koyu.png", optimize=True)
    sample = json.loads((DATA_OUT / "ornek-cekim-plani.json").read_text(encoding="utf-8"))
    sample = {
        "$kaynak": "Rast Creative Studio · https://rastcreative.com/araclar/cekim-plani/ · CC BY 4.0",
        "$not": "Örnek veridir. Araçta “JSON yükle” ile açın.",
        **sample,
    }
    (out / "rast-ornek-cekim-plani.json").write_text(json.dumps(sample, ensure_ascii=False, indent=2), encoding="utf-8")


def planner_sample(base: str):
    """Print the sample plan from the BUILT page (serve dist/ at `base`)."""
    out = PUB / "cekim-plani"
    raw = TMP / "plan.raw.pdf"
    url = base.rstrip("/") + "/araclar/cekim-plani/?ornek=1&baski=1"
    chrome(["--no-pdf-header-footer", "--virtual-time-budget=10000", "--run-all-compositor-stages-before-draw", f"--print-to-pdf={raw}", url])
    pdf = fitz.open(raw)
    pdf.set_metadata({
        "title": "Örnek çekim planı (call sheet)", "author": "Rast Creative Studio",
        "subject": "rastcreative.com/araclar/cekim-plani/ ile hazırlandı", "keywords": "call sheet, çekim planı, rastcreative.com",
        "creator": "Rast Creative Studio · rastcreative.com", "producer": "rastcreative.com",
    })
    dest = out / "rast-ornek-cekim-plani.pdf"
    pdf.save(dest, garbage=3, deflate=True)
    n = pdf.page_count
    pdf.close()
    pv = previews(dest, "rast-ornek-cekim-plani", list(range(1, n + 1)), 900)
    print("planner sample:", n, "pages", pv, "QA:", TMP)


def main():
    if len(sys.argv) > 2 and sys.argv[1] == "--plan-ornek":
        planner_sample(sys.argv[2])
        return
    planner_assets()
    report = {}
    DATA_OUT.mkdir(parents=True, exist_ok=True)

    # 1. storyboard + shot list
    sb = C.STORYBOARD
    out = PUB / sb["slug"]
    out.mkdir(parents=True, exist_ok=True)
    files = []
    for vertical, key in ((False, "sb169"), (True, "sb916")):
        d = storyboard_doc(vertical)
        p = out / f"{sb['files'][key]}.pdf"
        report[p.name] = build_pdf(d, p, {"title": d.title, "subject": "Storyboard şablonu", "keywords": "storyboard, şablon"})
        files.append(p)
    shot = HtmlDoc("Plan Listesi (Shot List)", "Çekim sırasına göre her planın ölçeği, açısı, hareketi ve notu. Arka sayfa: hızlı başvuru.", sb["slug"], True)
    for b in sb["shot_blocks"]:
        shot.block(b)
    p = out / f"{sb['files']['shot']}.pdf"
    report[p.name] = build_pdf(shot, p, {"title": "Plan Listesi (Shot List)", "subject": "Plan listesi şablonu", "keywords": "shot list, plan listesi"})
    files.append(p)
    pv = previews(out / "rast-storyboard-16x9.pdf", "rast-storyboard-16x9", [1], 1000)
    pv += previews(out / "rast-storyboard-9x16.pdf", "rast-storyboard-9x16", [1], 1000)
    pv += previews(p, "rast-plan-listesi", [1, 2], 1000)
    docx_p = out / "rast-plan-listesi.docx"
    docx_build("Plan Listesi (Shot List)", "Rast Creative Studio · ücretsiz şablon", sb["slug"], sb["shot_blocks"], docx_p, landscape=True)
    md_p = out / "rast-plan-listesi-notion.md"
    md_build("Plan Listesi (Shot List)", "Çekim sırasına göre her planın ölçeği, açısı, hareketi ve notu.", sb["slug"], sb["shot_blocks"], md_p)
    xl, cs = out / "rast-plan-listesi.xlsx", out / "rast-plan-listesi.csv"
    shotlist_xlsx(xl, cs)
    files += [docx_p, xl, cs, md_p]
    bundle(sb["slug"], "Storyboard ve Plan Listesi Şablonu", files)
    (DATA_OUT / f"{sb['slug']}.json").write_text(json.dumps({
        "previews": pv,
        "olcekler": C.OLCEKLER, "hareketler": C.HAREKETLER, "acilar": C.ACILAR,
        "shotCols": [h for h, _ in C.SHOT_COLS], "kurallar": C.SHOT_RULES,
    }, ensure_ascii=False, indent=1), encoding="utf-8")

    # 2-4. form-style documents
    for spec, pv_pages, extra in (
        (C.BRIEF, [1, 2, 3], ""),
        (C.CHECKLIST, [1, 2, 3], ""),
        (C.IZIN, [1, 2, 3], "\nÖNEMLİ: İzin formu genel bir taslaktır, hukuki danışmanlık değildir; kullanmadan önce bir avukata kontrol ettirin.\n"),
    ):
        slug, base = spec["slug"], spec["file"]
        out = PUB / slug
        out.mkdir(parents=True, exist_ok=True)
        d = HtmlDoc(spec["title"], spec["subtitle"], slug)
        for b in spec["blocks"]:
            d.block(b)
        p = out / f"{base}.pdf"
        report[p.name] = build_pdf(d, p, {"title": spec["title"], "subject": spec["subtitle"], "keywords": spec["title"]})
        pv = previews(p, base, pv_pages)
        docx_p, md_p = out / f"{base}.docx", out / f"{base}-notion.md"
        docx_build(spec["title"], spec["subtitle"], slug, spec["blocks"], docx_p)
        md_build(spec["title"], spec["subtitle"], slug, spec["blocks"], md_p)
        files = [p, docx_p, md_p]
        if spec is C.CHECKLIST:
            xl, cs = out / f"{base}.xlsx", out / f"{base}.csv"
            checklist_xlsx(xl, cs)
            files += [xl, cs]
        bundle(slug, spec["title"], files, extra)
        outline = []
        for b in spec["blocks"]:
            if b[0] == "h":
                outline.append({"h": b[1], "items": []})
            elif outline and b[0] in ("checks",):
                outline[-1]["items"].append(b[1])
            elif outline and b[0] in ("area",):
                outline[-1]["items"].append(b[1])
            elif outline and b[0] == "row":
                outline[-1]["items"] += [lb for lb, _n, _f in b[1]]
            elif outline and b[0] == "dl":
                outline[-1]["items"] += [a for a, _ in b[1]]
        dat = {"previews": pv, "outline": outline, "pages": len(d.pages)}
        if spec is C.CHECKLIST:
            dat["sections"] = [{"name": n, "items": [{"item": a, "tip": t} for a, t in its]} for n, its in C.CHECKLIST_SECTIONS]
        (DATA_OUT / f"{slug}.json").write_text(json.dumps(dat, ensure_ascii=False, indent=1), encoding="utf-8")

    print(json.dumps(report, ensure_ascii=False, indent=1))
    print("QA PNGs:", TMP)
    bad = {k: v["problems"] for k, v in report.items() if v["problems"]}
    if bad:
        print("PROBLEMS:", json.dumps(bad, ensure_ascii=False))
        sys.exit(2)


if __name__ == "__main__":
    main()
